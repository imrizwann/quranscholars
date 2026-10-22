"""Export frontend guide Word editions using a Python runtime with python-docx.

After export, render each DOCX with the document skill renderer and inspect every
page. Export alone does not verify visual layout.
"""
from pathlib import Path
import re
import sys
import textwrap
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def inline(paragraph, value):
    for token in re.split(r"(`[^`]+`)", value):
        run = paragraph.add_run(token.strip("`"))
        if token.startswith("`"):
            run.font.name = "Consolas"
            run.font.size = Pt(9)


def export(source):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.8)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet", "Footer"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(7)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.08
    doc.styles["Title"].font.size = Pt(25)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.add_run("QuranScholar | 8 October 2026 | Page ").font.size = Pt(9)
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)
    lines = source.read_text(encoding="utf-8").splitlines()
    doc.core_properties.title = lines[0][2:]
    doc.core_properties.subject = source.parent.name.replace("-", " ")
    headings = [line[3:] for line in lines if line.startswith("## ")]
    contents_written = False
    in_code = False
    for line in lines:
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run("\n".join(textwrap.wrap(line, 80, break_long_words=True, replace_whitespace=False)) or " ")
            run.font.name, run.font.size = "Consolas", Pt(8)
        elif line.startswith("# "):
            doc.add_paragraph(line[2:], "Title")
        elif line.startswith("## "):
            if not contents_written:
                doc.add_paragraph("Contents", "Heading 1")
                for heading in headings:
                    doc.add_paragraph(heading)
                doc.add_page_break()
                contents_written = True
            doc.add_paragraph(line[3:], "Heading 1")
        elif line.startswith("- "):
            inline(doc.add_paragraph(style="List Bullet"), line[2:])
        elif line.strip():
            inline(doc.add_paragraph(), line)
    output = source.with_suffix(".docx")
    doc.save(output)
    print(output)


if __name__ == "__main__":
    feature = sys.argv[1] if len(sys.argv) > 1 else "recorded-courses"
    for name in ("01-current-implementation.md", "02-next-implementation-guide.md"):
        export(ROOT / "docs" / feature / name)
